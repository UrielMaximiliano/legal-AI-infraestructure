"""Tests focales del parser DOCX read-only sin IA ni PII en logs."""

from __future__ import annotations

import hashlib
import io
import json
import logging
import zipfile
from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocumentObject
from lxml import etree

from legal_ai.application.docx_parser import (
    DOCX_MIME,
    DocxParseError,
    DocxParser,
    DocxParserConfig,
    DocxParseResult,
)
from legal_ai.domain.docx_placeholders import (
    FieldKind,
    PlaceholderCandidate,
    PlaceholderOrigin,
    PlaceholderStatus,
    PlaceholderSyntax,
    infer_field_kind,
    normalize_key,
)


def _save_bytes(document: DocumentObject) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _basic_doc_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Sin placeholders.")
    return _save_bytes(document)


def _parse_bytes(data: bytes, filename: str = "plantilla.docx") -> DocxParseResult:
    return DocxParser().parse_bytes(data, filename=filename, declared_mime=DOCX_MIME)


def _keys(result: DocxParseResult) -> set[str]:
    return {c.normalized_key for c in result.candidates}


def _by_key(result: DocxParseResult, key: str) -> PlaceholderCandidate:
    for candidate in result.candidates:
        if candidate.normalized_key == key:
            return candidate
    raise AssertionError(f"missing key {key}")


def _patch_document_xml(data: bytes, transform: Callable[[bytes], bytes]) -> bytes:
    source = io.BytesIO(data)
    out = io.BytesIO()
    with (
        zipfile.ZipFile(source, "r") as zin,
        zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zout,
    ):
        for item in zin.infolist():
            payload = zin.read(item.filename)
            if item.filename == "word/document.xml":
                payload = transform(payload)
            zout.writestr(item, payload)
    return out.getvalue()


class TestSyntaxCoverage:
    def test_double_brace(self) -> None:
        document = Document()
        document.add_paragraph("Expediente {{expediente}} fin.")
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "expediente")
        assert candidate.syntax == PlaceholderSyntax.DOUBLE_BRACE
        assert candidate.original_text == "{{expediente}}"
        assert candidate.confidence == pytest.approx(0.95)
        assert candidate.status == PlaceholderStatus.DETECTED
        assert candidate.explanation.strip() != ""

    def test_single_brace(self) -> None:
        document = Document()
        document.add_paragraph("Hola {nombre} fin.")
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "nombre")
        assert candidate.syntax == PlaceholderSyntax.SINGLE_BRACE

    def test_bracket(self) -> None:
        document = Document()
        document.add_paragraph("Firma [nombre_firmante] aquí.")
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "nombre_firmante")
        assert candidate.syntax == PlaceholderSyntax.BRACKET

    def test_double_angle(self) -> None:
        document = Document()
        document.add_paragraph("Monto <<monto_total>> pesos.")
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "monto_total")
        assert candidate.syntax == PlaceholderSyntax.DOUBLE_ANGLE
        assert candidate.field_kind == FieldKind.CURRENCY

    def test_completar(self) -> None:
        document = Document()
        document.add_paragraph("Falta [COMPLETAR] urgente.")
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "completar")
        assert candidate.syntax == PlaceholderSyntax.COMPLETAR

    def test_completar_with_label(self) -> None:
        document = Document()
        document.add_paragraph("Falta [COMPLETAR: domicilio] ya.")
        result = _parse_bytes(_save_bytes(document))
        assert "domicilio" in _keys(result)

    def test_blank_with_label(self) -> None:
        document = Document()
        document.add_paragraph("Nombre: ____________ fin.")
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "nombre")
        assert candidate.syntax == PlaceholderSyntax.BLANK

    def test_blank_without_label_uses_fallback(self) -> None:
        document = Document()
        document.add_paragraph("____________")
        result = _parse_bytes(_save_bytes(document))
        assert len(result.candidates) == 1
        assert result.candidates[0].normalized_key.startswith("campo_")

    def test_date_blanks_grouped(self) -> None:
        document = Document()
        document.add_paragraph("Fecha: ____/____/______ fin.")
        result = _parse_bytes(_save_bytes(document))
        assert len(result.candidates) == 1
        assert result.candidates[0].normalized_key == "fecha"

    def test_ellipsis(self) -> None:
        document = Document()
        document.add_paragraph("Domicilio: ... fin.")
        result = _parse_bytes(_save_bytes(document))
        assert any(c.syntax == PlaceholderSyntax.ELLIPSIS for c in result.candidates)

    def test_unicode_ellipsis(self) -> None:
        document = Document()
        document.add_paragraph("Observaciones: … fin.")
        result = _parse_bytes(_save_bytes(document))
        assert any(c.syntax == PlaceholderSyntax.ELLIPSIS for c in result.candidates)


class TestLocations:
    def test_paragraph_origin(self) -> None:
        document = Document()
        document.add_paragraph("Uno {{primero}}.")
        document.add_paragraph("Dos {{segundo}}.")
        result = _parse_bytes(_save_bytes(document))
        first = _by_key(result, "primero")
        assert first.origin == PlaceholderOrigin.PARAGRAPH
        assert first.locations[0].paragraph_index == 0

    def test_table_origin_with_coordinates(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Nombre {{nombre}}"
        table.cell(1, 1).text = "Fecha {{fecha}}"
        result = _parse_bytes(_save_bytes(document))
        nombre = _by_key(result, "nombre")
        assert nombre.origin == PlaceholderOrigin.TABLE_CELL
        assert nombre.locations[0].table_row == 0
        assert nombre.locations[0].table_col == 0
        fecha = _by_key(result, "fecha")
        assert fecha.locations[0].table_row == 1
        assert fecha.locations[0].table_col == 1

    def test_header_and_footer(self) -> None:
        document = Document()
        document.add_paragraph("Cuerpo {{cuerpo}}.")
        section = document.sections[0]
        section.header.paragraphs[0].text = "Encabezado {{encabezado}}."
        section.footer.paragraphs[0].text = "Pie {{pie}}."
        result = _parse_bytes(_save_bytes(document))
        assert _by_key(result, "encabezado").origin == PlaceholderOrigin.HEADER
        assert _by_key(result, "pie").origin == PlaceholderOrigin.FOOTER
        assert _by_key(result, "cuerpo").origin == PlaceholderOrigin.PARAGRAPH

    def test_split_runs_are_detected(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Hola {{")
        paragraph.add_run("nom")
        paragraph.add_run("bre}} fin.")
        result = _parse_bytes(_save_bytes(document))
        assert "nombre" in _keys(result)

    def test_split_runs_double_angle(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.add_run("Monto <<mon")
        paragraph.add_run("to_total>> fin.")
        result = _parse_bytes(_save_bytes(document))
        assert "monto_total" in _keys(result)

    def test_repetition_groups_locations(self) -> None:
        document = Document()
        document.add_paragraph("Uno {{expediente}}.")
        document.add_paragraph("Dos {{expediente}}.")
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Tres {{expediente}}."
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "expediente")
        assert candidate.occurrences == 3
        assert len(candidate.locations) == 3
        origins = {loc.origin for loc in candidate.locations}
        assert PlaceholderOrigin.PARAGRAPH in origins
        assert PlaceholderOrigin.TABLE_CELL in origins


class TestCandidateShape:
    def test_all_fields_present(self) -> None:
        document = Document()
        document.add_paragraph("Fecha {{fecha_emision}}.")
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "fecha_emision")
        assert candidate.original_text.strip() != ""
        assert candidate.normalized_key == "fecha_emision"
        assert isinstance(candidate.origin, PlaceholderOrigin)
        assert 0.0 < candidate.confidence <= 1.0
        assert candidate.explanation.strip() != ""
        assert candidate.status == PlaceholderStatus.DETECTED
        assert candidate.field_kind == FieldKind.DATE

    def test_field_kind_inference(self) -> None:
        assert infer_field_kind("fecha_nac", PlaceholderSyntax.DOUBLE_BRACE) == (
            FieldKind.DATE
        )
        assert infer_field_kind("monto_total", PlaceholderSyntax.DOUBLE_BRACE) == (
            FieldKind.CURRENCY
        )
        assert infer_field_kind("email_contacto", PlaceholderSyntax.DOUBLE_BRACE) == (
            FieldKind.EMAIL
        )
        assert infer_field_kind("expediente", PlaceholderSyntax.DOUBLE_BRACE) == (
            FieldKind.NUMBER
        )
        assert infer_field_kind("observaciones", PlaceholderSyntax.DOUBLE_BRACE) == (
            FieldKind.TEXT
        )

    def test_normalize_key_deterministic(self) -> None:
        assert normalize_key("Nombre Completo") == "nombre_completo"
        assert normalize_key("  Monto-Total ") == "monto_total"
        assert normalize_key("!!!", fallback="campo_1") == "campo_1"

    def test_human_workflow_transitions(self) -> None:
        document = Document()
        document.add_paragraph("Hola {{nombre}}.")
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "nombre")
        candidate.request_review()
        assert candidate.status == PlaceholderStatus.PENDING_REVIEW
        candidate.resolve(confirmed=True, actor="revisor-humano")
        assert candidate.status == PlaceholderStatus.CONFIRMED

    def test_ai_actor_is_rejected(self) -> None:
        document = Document()
        document.add_paragraph("Hola {{nombre}}.")
        result = _parse_bytes(_save_bytes(document))
        candidate = _by_key(result, "nombre")
        candidate.request_review()
        with pytest.raises(ValueError, match="AI_CONFIRM"):
            candidate.resolve(confirmed=True, actor="ai-agent-01")


class TestNativeControls:
    def test_sdt_alias_is_detected(self) -> None:
        document = Document()
        document.add_paragraph("Antes.")
        data = _save_bytes(document)

        def _inject(payload: bytes) -> bytes:
            root = etree.fromstring(payload)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            body = root.find("w:body", ns)
            assert body is not None
            sdt_xml = (
                '<w:sdt xmlns:w="http://schemas.openxmlformats.org'
                '/wordprocessingml/2006/main">'
                '<w:sdtPr><w:alias w:val="nombre_sdt"/>'
                '<w:tag w:val="nombre_sdt"/></w:sdtPr>'
                "<w:sdtContent><w:p><w:r><w:t>Valor</w:t></w:r></w:p>"
                "</w:sdtContent></w:sdt>"
            )
            sdt = etree.fromstring(sdt_xml.encode("utf-8"))
            sect = body.find("w:sectPr", ns)
            if sect is not None:
                sect.addprevious(sdt)
            else:
                body.append(sdt)
            return etree.tostring(root, xml_declaration=True, encoding="UTF-8")

        patched = _patch_document_xml(data, _inject)
        result = _parse_bytes(patched)
        candidate = _by_key(result, "nombre_sdt")
        assert candidate.syntax == PlaceholderSyntax.NATIVE_SDT

    def test_mergefield_is_detected(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        element = paragraph._element
        for child in list(element):
            element.remove(child)
        ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        parts: tuple[tuple[str, str], ...] = (
            ("fldChar_begin", ""),
            ("instr", " MERGEFIELD apellido "),
            ("fldChar_end", ""),
        )
        for tag, text in parts:
            run = etree.SubElement(element, f"{{{ns}}}r")
            if tag.startswith("fldChar"):
                fld = etree.SubElement(run, f"{{{ns}}}fldChar")
                kind = "begin" if "begin" in tag else "end"
                fld.set(f"{{{ns}}}fldCharType", kind)
            else:
                instr = etree.SubElement(run, f"{{{ns}}}instrText")
                instr.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                instr.text = text
        data = _save_bytes(document)
        result = _parse_bytes(data)
        candidate = _by_key(result, "apellido")
        assert candidate.syntax == PlaceholderSyntax.NATIVE_MERGEFIELD


class TestSecurity:
    def test_extension_docm_rejected(self) -> None:
        with pytest.raises(DocxParseError, match="MACRO|EXTENSION"):
            DocxParser().parse_bytes(b"fake", filename="plantilla.docm")

    def test_mime_mismatch_rejected(self) -> None:
        with pytest.raises(DocxParseError, match="MIME"):
            DocxParser().parse_bytes(
                _basic_doc_bytes(),
                filename="plantilla.docx",
                declared_mime="application/pdf",
            )

    def test_corrupt_rejected(self, tmp_path: Path) -> None:
        path = tmp_path / "roto.docx"
        path.write_bytes(b"no es un zip")
        with pytest.raises(DocxParseError, match="CORRUPT|STRUCTURE"):
            DocxParser().parse_path(path, declared_mime=DOCX_MIME)

    def test_vba_rejected(self) -> None:
        data = _basic_doc_bytes()
        buffer = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(data), "r") as zin,
            zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as zout,
        ):
            for item in zin.infolist():
                zout.writestr(item, zin.read(item.filename))
            zout.writestr("word/vbaProject.bin", b"macro")
        with pytest.raises(DocxParseError, match="MACRO"):
            _parse_bytes(buffer.getvalue())

    def test_external_rel_rejected(self) -> None:
        data = _basic_doc_bytes()

        def _inject(payload: bytes) -> bytes:
            text = payload.decode("utf-8")
            rel_type = (
                "http://schemas.openxmlformats.org/officeDocument/2006"
                "/relationships/oleObject"
            )
            return text.replace(
                "</Relationships>",
                '<Relationship Id="rId999" Type="'
                + rel_type
                + '" Target="http://externo.example/objeto" '
                + 'TargetMode="External"/></Relationships>',
            ).encode("utf-8")

        out = io.BytesIO()
        with (
            zipfile.ZipFile(io.BytesIO(data), "r") as zin,
            zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zout,
        ):
            for item in zin.infolist():
                payload = zin.read(item.filename)
                if item.filename.endswith(".rels"):
                    payload = _inject(payload)
                zout.writestr(item, payload)
        with pytest.raises(DocxParseError, match="EXTERNAL"):
            _parse_bytes(out.getvalue())

    def test_path_traversal_rejected(self) -> None:
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_STORED) as archive:
            archive.writestr("[Content_Types].xml", b"x")
            archive.writestr("word/document.xml", b"x")
            archive.writestr("../evil.xml", b"x")
        with pytest.raises(DocxParseError, match="TRAVERSAL|CORRUPT|STRUCTURE"):
            _parse_bytes(buffer.getvalue())

    def test_doctype_rejected(self) -> None:
        data = _basic_doc_bytes()
        source = io.BytesIO(data)
        out = io.BytesIO()
        with (
            zipfile.ZipFile(source, "r") as zin,
            zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zout,
        ):
            for item in zin.infolist():
                payload = zin.read(item.filename)
                if item.filename == "word/document.xml":
                    prefix = b'<?xml version="1.0"?><!DOCTYPE foo [<!ENTITY x "y">]>'
                    payload = prefix + payload
                zout.writestr(item, payload)
        with pytest.raises(DocxParseError, match="XML|CORRUPT"):
            _parse_bytes(out.getvalue())

    def test_limits_enforced(self) -> None:
        config = DocxParserConfig(
            max_file_bytes=10,
            max_uncompressed_bytes=10,
            max_entries=500,
            max_compression_ratio=100,
            max_candidates=1000,
            max_text_bytes=2 * 1024 * 1024,
        )
        with pytest.raises(DocxParseError, match="SIZE|LIMIT"):
            DocxParser(config).parse_bytes(
                _basic_doc_bytes(), filename="plantilla.docx"
            )

    def test_too_many_fields_rejected(self) -> None:
        document = Document()
        document.add_paragraph(" ".join(f"{{{{campo{i}}}}}" for i in range(5)))
        data = _save_bytes(document)
        config = DocxParserConfig(
            max_file_bytes=20 * 1024 * 1024,
            max_uncompressed_bytes=50 * 1024 * 1024,
            max_entries=500,
            max_compression_ratio=100,
            max_candidates=2,
            max_text_bytes=2 * 1024 * 1024,
        )
        with pytest.raises(DocxParseError, match="TOO_MANY"):
            DocxParser(config).parse_bytes(data, filename="plantilla.docx")

    def test_original_not_modified(self, tmp_path: Path) -> None:
        document = Document()
        document.add_paragraph("Hola {{nombre}}.")
        path = tmp_path / "original.docx"
        path.write_bytes(_save_bytes(document))
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        mtime_before = path.stat().st_mtime_ns
        DocxParser().parse_path(path, declared_mime=DOCX_MIME)
        assert hashlib.sha256(path.read_bytes()).hexdigest() == before
        assert path.stat().st_mtime_ns == mtime_before

    def test_logs_do_not_contain_content_or_pii(
        self, caplog: pytest.LogCaptureFixture
    ) -> None:
        document = Document()
        document.add_paragraph("DNI {{dni_persona}} secreto.")
        data = _save_bytes(document)
        with caplog.at_level(logging.INFO, logger="legal_ai.004"):
            _parse_bytes(data)
        rendered = "\n".join(
            json.dumps(record.__dict__, default=str) for record in caplog.records
        )
        assert "dni_persona" not in rendered
        assert "{{" not in rendered
        assert "secreto" not in rendered.lower()
