"""Focused checks for configurable template parsing and preview."""

from legal_ai.application.template_management import (
    extract_file_content,
    fields_from_markers,
    marker_keys,
    normalize_blank_fields,
    preview_document,
    suggestions_for_body,
    validate_values,
)


def test_repeated_markers_create_one_field_in_first_seen_order() -> None:
    body = "Resolución {{expediente}}\n{{interesado}} y {{expediente}}"

    assert marker_keys(body) == ["expediente", "interesado"]
    assert [field["key"] for field in fields_from_markers(body)] == [
        "expediente",
        "interesado",
    ]


def test_preview_preserves_false_and_zero_values() -> None:
    document, errors, _warnings = preview_document(
        name="Prueba",
        document_type="resolucion",
        body="Monto: {{monto}}; Activo: {{activo}}",
        fields=[
            {"key": "monto", "label": "Monto", "type": "number"},
            {"key": "activo", "label": "Activo", "type": "boolean"},
        ],
        rules=[],
        blocks=None,
        values={"monto": "0", "activo": "false"},
    )

    assert errors == []
    assert "Monto: 0" in document["content"]
    assert "Activo: False" in document["content"]


def test_placeholder_suggestions_group_occurrences() -> None:
    suggestions = suggestions_for_body(
        "Firmante: ____________\nRepite: ____________", fields=[]
    )

    assert len(suggestions) == 1
    assert len(suggestions[0]["occurrences"]) == 2


def test_blanks_use_line_label() -> None:
    assert normalize_blank_fields("Nombre: ____________") == "Nombre: {{nombre}}"
    assert normalize_blank_fields("Monto: $ ____________") == "Monto: $ {{monto}}"


def test_trailing_spaces_after_line_label_become_field() -> None:
    assert normalize_blank_fields("Nombre:             ") == "Nombre: {{nombre}}"
    assert normalize_blank_fields("Texto normal con espacios    ") == (
        "Texto normal con espacios    "
    )


def test_date_segments_group_into_single_marker() -> None:
    result = normalize_blank_fields("Fecha: ____/____/______")

    assert result == "Fecha: {{fecha}}"
    assert marker_keys(result) == ["fecha"]


def test_duplicate_labels_reuse_same_key() -> None:
    result = normalize_blank_fields("Nombre: ____________\nNombre: ____________")

    assert result == "Nombre: {{nombre}}\nNombre: {{nombre}}"
    assert marker_keys(result) == ["nombre"]


def test_fallback_campo_n_without_label_is_stable() -> None:
    result = normalize_blank_fields("____________\n____________")

    assert result == "{{campo_1}}\n{{campo_2}}"


def test_preserves_markers_and_does_not_convert_normal_text() -> None:
    body = "Expediente {{expediente}}.\nTexto normal sin blancos."
    assert normalize_blank_fields(body) == body
    assert normalize_blank_fields("Referencia [Art. 1] del Código Civil.") == (
        "Referencia [Art. 1] del Código Civil."
    )

    preserved = normalize_blank_fields("Folio {{folio}}: ____________")
    assert "{{folio}}" in preserved
    assert "____________" not in preserved


def test_extract_docx_normalizes_body_and_blocks() -> None:
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_paragraph("Nombre: ____________")
    payload = BytesIO()
    document.save(payload)

    body, blocks, _warnings, _pages = extract_file_content(
        "plantilla.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        payload.getvalue(),
    )

    assert body == "Nombre: {{nombre}}"
    assert blocks[0]["content"] == "Nombre: {{nombre}}"


def test_extract_docx_includes_header_footer_and_empty_table_fields() -> None:
    from io import BytesIO

    from docx import Document

    document = Document()
    document.add_paragraph("Cuerpo")
    document.add_table(rows=1, cols=2)
    section = document.sections[0]
    section.header.paragraphs[0].text = "Encabezado: ____________"
    section.footer.paragraphs[0].text = "Pie: ____________"
    payload = BytesIO()
    document.save(payload)

    body, _blocks, _warnings, _pages = extract_file_content(
        "plantilla.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        payload.getvalue(),
    )

    assert "Cuerpo" in body
    assert "Encabezado: {{encabezado}}" in body
    assert "Pie: {{pie}}" in body
    assert "{{celda_vacia_1_1}}" in body


def test_cuit_validation_checks_the_check_digit() -> None:
    fields = [{"key": "cuit", "type": "text"}]
    rules = [
        {
            "type": "validation",
            "field": "cuit",
            "format": "cuit",
            "message": "CUIT inválido",
        }
    ]

    assert validate_values(fields, rules, {"cuit": "20-32964233-0"}) == []
    assert validate_values(fields, rules, {"cuit": "20-32964233-1"}) == [
        "CUIT inválido"
    ]
